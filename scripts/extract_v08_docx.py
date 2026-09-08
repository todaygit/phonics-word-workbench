from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def main() -> None:
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    document = Document(source)
    blocks: list[dict[str, object]] = []

    for index, block in enumerate(iter_blocks(document)):
        if isinstance(block, Paragraph):
            text = "".join(run.text for run in block.runs).strip()
            if text:
                blocks.append(
                    {
                        "index": index,
                        "type": "paragraph",
                        "style": block.style.name if block.style else "",
                        "text": text,
                    }
                )
        else:
            rows = []
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    text = "\n".join(
                        paragraph.text.strip()
                        for paragraph in cell.paragraphs
                        if paragraph.text.strip()
                    )
                    cells.append(text)
                rows.append(cells)
            blocks.append({"index": index, "type": "table", "rows": rows})

    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(
            {
                "source": str(source),
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
                "blocks": blocks,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
