import json
import re
from pathlib import Path

from docx import Document


SOURCE = Path(r"C:\Users\Today_ztt\Desktop\自然拼读背单词小手册_拼写导向版_v1.9.docx")
OUT = Path(__file__).resolve().parents[1] / "app" / "word-bank-v19.json"


def cells(table):
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def main():
    doc = Document(SOURCE)
    rule_rows = cells(doc.tables[3])[1:]
    titles = {}
    paragraphs = [p.text.strip() for p in doc.paragraphs]
    for index, text in enumerate(paragraphs):
        match = re.match(r"^(\d+\.\d+)\s+(.+)$", text)
        if not match:
            continue
        chapter_id = match.group(1)
        rule = ""
        child_note = ""
        for j in range(index + 1, min(index + 14, len(paragraphs))):
            if paragraphs[j] in {"一、本课规则", "二、给孩子的讲法"}:
                continue
            if paragraphs[j] in {"三、词库与句型练习表", "四、", ""}:
                continue
            if not rule:
                rule = paragraphs[j]
            elif not child_note:
                child_note = paragraphs[j]
                break
        titles[chapter_id] = {"title": text, "rule": rule, "childNote": child_note}

    chapters = []
    words = []
    for order, row in enumerate(rule_rows, start=1):
        match = re.match(r"^(\d+\.\d+)\s+(.+)$", row[0])
        chapter_id = match.group(1) if match else row[0]
        title = match.group(2) if match else row[0]
        rule_desc = row[1]
        meta = titles.get(chapter_id, {})
        chapter_title = meta.get("title") or f"{chapter_id} {title}"
        chapter_rule = meta.get("rule") or rule_desc
        chapter_note = meta.get("childNote") or "先听音、看拆分，再尝试拼写。"
        table_rows = cells(doc.tables[3 + order])
        chapter_words = []
        for seq, word_row in enumerate(table_rows[1:], start=1):
            if len(word_row) < 7 or not word_row[2]:
                continue
            item = {
                "id": f"v19-{order:03d}-{seq:03d}",
                "chapterId": chapter_id,
                "chapterOrder": order,
                "wordOrder": seq,
                "level": "v1.9",
                "word": word_row[2],
                "ipa": word_row[4],
                "phonics": word_row[5],
                "meaning": word_row[1],
                "example": word_row[6],
                "partOfSpeech": word_row[3],
            }
            words.append(item)
            chapter_words.append(item)
        chapters.append(
            {
                "id": chapter_id,
                "order": order,
                "title": chapter_title,
                "rule": chapter_rule,
                "childNote": chapter_note,
                "wordCount": len(chapter_words),
            }
        )

    payload = {
        "version": "v1.9",
        "sourceTitle": "自然拼读背单词小手册·拼写导向版 v1.9",
        "chapters": chapters,
        "words": words,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"wrote {len(chapters)} chapters and {len(words)} words to {OUT}")


if __name__ == "__main__":
    main()
